"""Line breaks in a translation: one original paragraph may be several Ukrainian ones (a dialogue).
Every way a translation is written keeps LF; the .docx export makes a paragraph of every line."""
import tempfile
import unittest
from io import BytesIO
from pathlib import Path

from docx import Document

from backend.export_service import ExportService
from backend.storage import Storage, normalize_translation_text
from backend.translations.service import TranslationService

DIALOGUE = "— Привіт, — сказав він.\n— Привіт, — відповіла вона."


class NormalizeTests(unittest.TestCase):
    def test_breaks_are_kept_and_tidied(self):
        self.assertEqual("а\nб", normalize_translation_text("а\nб"))
        self.assertEqual("а\nб", normalize_translation_text("  а  \r\n\r\n\n\t б \r"))
        self.assertEqual("а\nб", normalize_translation_text("\n\nа\n\n\n\nб\n\n"))
        self.assertEqual("", normalize_translation_text("\n \n"))

    def test_text_without_a_break_is_left_alone(self):
        for text in (None, "", "  Дрейвен  прийшов  ", "<b>а</b> б"):
            self.assertEqual(text, normalize_translation_text(text))


class Fixture(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.storage = Storage(Path(self.directory.name) / "w.sqlite3")
        self.project_id = self.storage.create_project({"title": "LB", "status": "translation"})["projectId"]
        self.storage.save_book_structure(self.project_id, "b.epub", "application/epub+zip", b"b", {"chapters": [
            {"title": "One", "elements": [{"type": "paragraph", "text": "He said hi. She said hi."},
                                           {"type": "paragraph", "text": "Second"}]},
        ]})
        chapter = self.storage.get_book_structure(self.project_id)["chapters"][0]
        self.chapter_id = chapter["chapterId"]
        self.first, self.second = (element["paragraphId"] for element in chapter["elements"])

    def tearDown(self):
        self.directory.cleanup()

    def paragraph(self, paragraph_id=None):
        paragraph_id = paragraph_id or self.first
        for element in self.storage.get_book_structure(self.project_id)["chapters"][0]["elements"]:
            if element["paragraphId"] == paragraph_id:
                return element


class SavePathTests(Fixture):
    def test_editor_save_keeps_breaks_in_text_and_manual_variant(self):
        self.storage.save_paragraph_edit(self.first, "  " + DIALOGUE.replace("\n", "  \n\n\n  "), False)
        paragraph = self.paragraph()
        self.assertEqual(DIALOGUE, paragraph["translationText"])
        manual = [v for v in paragraph["translationVariants"] if v["engine"] == "manual"]
        self.assertEqual([DIALOGUE], [v["text"] for v in manual])
        # saving the same text again records nothing new
        self.storage.save_paragraph_edit(self.first, DIALOGUE, True)
        self.assertEqual(1, len([v for v in self.paragraph()["translationVariants"] if v["engine"] == "manual"]))

    def test_low_level_update_keeps_breaks(self):
        self.storage.update_paragraph(self.first, "а\r\nб", False)
        self.assertEqual("а\nб", self.paragraph()["translationText"])

    def test_engine_variant_keeps_breaks_and_is_chosen(self):
        self.storage.add_translation_variant(self.first, "claude", "— Так.\n\n— Ні.", "m", None)
        paragraph = self.paragraph()
        self.assertEqual("— Так.\n— Ні.", paragraph["translationText"])
        self.assertEqual(["— Так.\n— Ні."], [v["text"] for v in paragraph["translationVariants"] if v["engine"] == "claude"])

    def test_engine_answer_with_a_break_inside_p_reaches_the_variant(self):
        xml = f'<chunk><p id="{self.first}">\n  — Так.\n— Ні.\n</p><p id="{self.second}">Один</p></chunk>'
        mapped = TranslationService._parse_chunk_xml_result(xml, [self.first, self.second])
        self.assertIn("— Так.\n— Ні.", mapped[self.first])
        self.storage.add_translation_variant(self.first, "claude", mapped[self.first], "m", None)
        self.assertEqual("— Так.\n— Ні.", self.paragraph()["translationText"])

    def test_find_and_replace_keeps_breaks_and_undo_redo_restore_them(self):
        self.storage.save_paragraph_edit(self.first, "— Дрейвен прийшов.\n— Дрейвена бачили.", False)
        request = {"query": "Дрейвен", "replacement": "Дравен", "options": {"mode": "start", "matchCase": False},
                   "scope": "project", "chapterId": self.chapter_id}
        done = self.storage.replace_in_translation(self.project_id, request)
        self.assertEqual("— Дравен прийшов.\n— Дравена бачили.", self.paragraph()["translationText"])
        self.storage.restore_replacement(self.project_id, done["undo"]["items"], self.chapter_id)
        self.assertEqual("— Дрейвен прийшов.\n— Дрейвена бачили.", self.paragraph()["translationText"])
        items = [{k: v for k, v in item.items() if k != "previous"} for item in done["undo"]["items"]]
        self.storage.apply_replacement(self.project_id, items, None)
        self.assertEqual("— Дравен прийшов.\n— Дравена бачили.", self.paragraph()["translationText"])

    def test_binding_to_the_original_paragraph_is_unchanged(self):
        self.storage.save_paragraph_edit(self.first, DIALOGUE, False)
        paragraph = self.paragraph()
        self.assertEqual(self.first, paragraph["paragraphId"])
        self.assertEqual("He said hi. She said hi.", paragraph["originalText"])
        self.assertEqual(2, len(self.storage.get_book_structure(self.project_id)["chapters"][0]["elements"]))


class DocxTests(Fixture):
    def export(self, fmt):
        return Document(BytesIO(ExportService(self.storage).generate_docx(self.project_id, fmt).getvalue()))

    def test_translation_only_makes_a_paragraph_of_every_line_with_the_same_indent(self):
        self.storage.save_paragraph_edit(self.first, DIALOGUE + "\n— Втретє.", False)
        self.storage.save_paragraph_edit(self.second, "Звичайний абзац", False)
        document = self.export("translation_only")
        paragraphs = document.paragraphs
        self.assertEqual(["One", "— Привіт, — сказав він.", "— Привіт, — відповіла вона.", "— Втретє.", "Звичайний абзац"],
                         [p.text for p in paragraphs])
        self.assertIsNone(paragraphs[0].paragraph_format.first_line_indent)  # heading
        for paragraph in paragraphs[1:]:
            self.assertAlmostEqual(1.0, paragraph.paragraph_format.first_line_indent.cm, places=2)
            self.assertNotIn("<w:br", paragraph._p.xml)  # a real paragraph, not a soft break
        section = document.sections[0]
        for margin in (section.top_margin, section.bottom_margin, section.left_margin, section.right_margin):
            self.assertAlmostEqual(0.6, margin.cm, places=2)

    def test_bilingual_has_several_paragraphs_in_one_cell_and_no_indent(self):
        self.storage.save_paragraph_edit(self.first, DIALOGUE, False)
        document = self.export("bilingual")
        cell = document.tables[0].rows[1].cells[1]
        self.assertEqual(["— Привіт, — сказав він.", "— Привіт, — відповіла вона."], [p.text for p in cell.paragraphs])
        self.assertEqual("He said hi. She said hi.", document.tables[0].rows[1].cells[0].text)
        for paragraph in cell.paragraphs + document.paragraphs:
            self.assertIsNone(paragraph.paragraph_format.first_line_indent)
        self.assertAlmostEqual(0.6, document.sections[0].left_margin.cm, places=2)

    def test_inline_formatting_spanning_a_break_is_closed_and_reopened(self):
        self.storage.save_paragraph_edit(self.first, "<i>перший\nдругий</i> кінець", False)
        paragraphs = self.export("translation_only").paragraphs[1:3]
        self.assertEqual(["перший", "другий кінець"], [p.text for p in paragraphs])
        self.assertTrue(paragraphs[0].runs[0].italic)
        self.assertTrue(paragraphs[1].runs[0].italic)
        self.assertFalse(paragraphs[1].runs[1].italic)

    def test_text_without_breaks_gives_one_paragraph_each(self):
        self.storage.save_paragraph_edit(self.first, "Один абзац", False)
        self.assertEqual(["One", "Один абзац"], [p.text for p in self.export("translation_only").paragraphs][:2])
        self.assertEqual(1, len(self.export("bilingual").tables[0].rows[1].cells[1].paragraphs))


if __name__ == "__main__":
    unittest.main()
