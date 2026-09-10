import base64
import tempfile
import unittest
from io import BytesIO
from pathlib import Path

from docx import Document
from PIL import Image

from backend.export_service import ExportService
from backend.storage import Storage


def _sample_image_base64() -> str:
    buffer = BytesIO()
    Image.new("RGB", (4, 4), color="red").save(buffer, format="PNG")
    return base64.b64encode(buffer.getvalue()).decode("ascii")


class ExportServiceTests(unittest.TestCase):
    def setUp(self):
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.storage = Storage(Path(self.temporary_directory.name) / "workbench.sqlite3")
        project = self.storage.create_project({"title": "Export book", "status": "translation"})
        self.project_id = project["projectId"]
        structure = self.storage.save_book_structure(
            self.project_id,
            "book.epub",
            "application/epub+zip",
            b"book",
            {
                "chapters": [
                    {
                        "title": "Original chapter",
                        "elements": [
                            {"type": "paragraph", "text": "Original one"},
                            {"type": "paragraph", "text": "Service original"},
                        ],
                    },
                    {"title": "Excluded chapter", "elements": [{"type": "paragraph", "text": "Hidden original"}]},
                ]
            },
        )
        first, excluded = structure["chapters"]
        self.storage.update_chapter_title(first["chapterId"], "Перекладений розділ", True)
        self.storage.update_paragraph(first["elements"][0]["paragraphId"], "Переклад один", True)
        self.storage.update_paragraph(first["elements"][1]["paragraphId"], "Службовий переклад", True, True)
        self.storage.update_paragraph(excluded["elements"][0]["paragraphId"], "Прихований переклад", True)
        self.storage.set_chapter_export_flag(excluded["chapterId"], True)
        self.service = ExportService(self.storage)

    def tearDown(self):
        self.temporary_directory.cleanup()

    def test_generates_bilingual_table_and_includes_service_paragraphs(self):
        document = Document(self.service.generate_docx(self.project_id, "bilingual"))

        self.assertEqual("Перекладений розділ", document.paragraphs[0].text)
        self.assertEqual(1, len(document.tables))
        rows = [[cell.text for cell in row.cells] for row in document.tables[0].rows]
        self.assertEqual(
            [
                ["Оригінал", "Переклад"],
                ["Original one", "Переклад один"],
                ["Service original", "Службовий переклад"],
            ],
            rows,
        )
        self.assertNotIn("Hidden original", str(rows))

    def test_generates_translation_only_paragraphs_and_excludes_chapter(self):
        document = Document(self.service.generate_docx(self.project_id, "translation_only"))
        text = [paragraph.text for paragraph in document.paragraphs]

        self.assertEqual(["Перекладений розділ", "Переклад один", "Службовий переклад"], text)
        self.assertNotIn("Прихований переклад", text)

    def test_preserves_nested_formatting_in_bilingual_runs(self):
        marked_text = "Before <b>bold <i>bold italic</i></b> <s>deleted</s> after"
        paragraph_id = self.storage.get_book_structure(self.project_id)["chapters"][0]["elements"][0]["paragraphId"]
        self.storage.update_paragraph(paragraph_id, marked_text, True)

        document = Document(self.service.generate_docx(self.project_id, "bilingual"))
        runs = document.tables[0].rows[1].cells[1].paragraphs[0].runs

        self.assertEqual(marked_text.replace("<b>", "").replace("</b>", "").replace("<i>", "").replace("</i>", "").replace("<s>", "").replace("</s>", ""), document.tables[0].rows[1].cells[1].text)
        self.assertEqual(
            ["Before ", "bold ", "bold italic", " ", "deleted", " after"],
            [run.text for run in runs],
        )
        self.assertFalse(runs[0].bold)
        self.assertTrue(runs[1].bold)
        self.assertTrue(runs[2].bold)
        self.assertTrue(runs[2].italic)
        self.assertTrue(runs[4].font.strike)

    def test_preserves_nested_formatting_in_translation_only_runs(self):
        marked_text = "<i>Intro <b>emphasis</b></i> and <s>removed</s>"
        paragraph_id = self.storage.get_book_structure(self.project_id)["chapters"][0]["elements"][0]["paragraphId"]
        self.storage.update_paragraph(paragraph_id, marked_text, True)

        document = Document(self.service.generate_docx(self.project_id, "translation_only"))
        paragraph = document.paragraphs[1]

        self.assertEqual("Intro emphasis and removed", paragraph.text)
        self.assertEqual(["Intro ", "emphasis", " and ", "removed"], [run.text for run in paragraph.runs])
        self.assertTrue(paragraph.runs[0].italic)
        self.assertTrue(paragraph.runs[1].bold)
        self.assertTrue(paragraph.runs[1].italic)
        self.assertTrue(paragraph.runs[3].font.strike)

    def test_literal_formatting_markers_round_trip_through_paragraph_update(self):
        marked_text = "Зміна <b>жирного</b>, <i>курсивного</i> і <s>закресленого</s>."
        paragraph_id = self.storage.get_book_structure(self.project_id)["chapters"][0]["elements"][0]["paragraphId"]

        self.storage.update_paragraph(paragraph_id, marked_text, False)

        saved = self.storage.get_book_structure(self.project_id)["chapters"][0]["elements"][0]["translationText"]
        self.assertEqual(marked_text, saved)

    def test_rejects_unknown_format(self):
        with self.assertRaisesRegex(ValueError, "Unsupported"):
            self.service.generate_docx(self.project_id, "pdf")


class ExportServiceImageTests(unittest.TestCase):
    def setUp(self):
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.storage = Storage(Path(self.temporary_directory.name) / "workbench.sqlite3")
        project = self.storage.create_project({"title": "Illustrated book", "status": "translation"})
        self.project_id = project["projectId"]
        structure = self.storage.save_book_structure(
            self.project_id,
            "book.epub",
            "application/epub+zip",
            b"book",
            {
                "chapters": [
                    {
                        "title": "Chapter with picture",
                        "elements": [
                            {"type": "paragraph", "text": "Before image"},
                            {"type": "image", "imageData": _sample_image_base64()},
                            {"type": "paragraph", "text": "After image"},
                        ],
                    },
                ]
            },
        )
        chapter = structure["chapters"][0]
        paragraph_elements = [element for element in chapter["elements"] if element["type"] == "paragraph"]
        self.storage.update_paragraph(paragraph_elements[0]["paragraphId"], "Перед зображенням", True)
        self.storage.update_paragraph(paragraph_elements[1]["paragraphId"], "Після зображення", True)
        self.service = ExportService(self.storage)

    def tearDown(self):
        self.temporary_directory.cleanup()

    def test_translation_only_inserts_image_inline_between_paragraphs(self):
        document = Document(self.service.generate_docx(self.project_id, "translation_only"))
        text_blocks = [paragraph.text for paragraph in document.paragraphs if paragraph.text]

        self.assertEqual(["Chapter with picture", "Перед зображенням", "Після зображення"], text_blocks)
        self.assertEqual(1, len(document.inline_shapes))

    def test_bilingual_marks_image_placement_and_appends_it_at_the_end(self):
        document = Document(self.service.generate_docx(self.project_id, "bilingual"))
        rows = [[cell.text for cell in row.cells] for row in document.tables[0].rows]

        self.assertEqual(
            [
                ["Оригінал", "Переклад"],
                ["Before image", "Перед зображенням"],
                ["[Зображення 1]", "[Зображення 1]"],
                ["After image", "Після зображення"],
            ],
            rows,
        )
        heading_texts = [paragraph.text for paragraph in document.paragraphs]
        self.assertIn("Зображення", heading_texts)
        self.assertIn("Зображення 1 ⬇️", heading_texts)
        self.assertEqual(1, len(document.inline_shapes))

    def test_get_image_as_png_returns_none_for_missing_image(self):
        self.assertIsNone(self.service._get_image_as_png("missing-image"))


if __name__ == "__main__":
    unittest.main()